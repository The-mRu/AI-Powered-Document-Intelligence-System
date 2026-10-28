"""Fixtures generate every test file at runtime, so the repo stays small and the
tests document exactly what each input looks like."""
import io
import os

import numpy as np
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFilter, ImageFont

LINES = ["INVOICE", "Invoice No: INV-2024-0042", "Supplier: Acme Trading Ltd",
         "Customer: Globex Corporation", "Subtotal: 1000.00", "Tax: 150.00", "Total: 1150.00 USD"]


def _font(size):
    p = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    return ImageFont.truetype(p, size) if os.path.exists(p) else ImageFont.load_default(size)


def make_text_image(lines=LINES, size=(1240, 1754), font_size=44) -> Image.Image:
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    y = 120
    for ln in lines:
        d.text((100, y), ln, fill="black", font=_font(font_size))
        y += font_size * 2
    return img


def degrade(img: Image.Image, angle=3, blur=1.4, noise=18) -> Image.Image:
    img = img.rotate(angle, expand=True, fillcolor="white").filter(ImageFilter.GaussianBlur(blur))
    arr = np.array(img).astype(np.int16)
    arr += np.random.default_rng(0).normal(0, noise, arr.shape).astype(np.int16)
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


@pytest.fixture
def digital_pdf(tmp_path):
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 100), f"Digital page {i + 1}. Invoice INV-2024-0042 total 1150.00 USD", fontsize=12)
    p = tmp_path / "digital.pdf"
    doc.save(p)
    return p


@pytest.fixture
def scanned_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    buf = io.BytesIO()
    make_text_image().save(buf, "PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    p = tmp_path / "scan.pdf"
    doc.save(p)
    return p


@pytest.fixture
def mixed_pdf(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Page one has a real text layer with enough characters.", fontsize=12)
    page = doc.new_page()
    buf = io.BytesIO()
    make_text_image().save(buf, "PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    p = tmp_path / "mixed.pdf"
    doc.save(p)
    return p


@pytest.fixture
def png_path(tmp_path):
    p = tmp_path / "clean.png"
    make_text_image().save(p)
    return p


@pytest.fixture
def poor_quality_jpg(tmp_path):
    p = tmp_path / "poor.jpg"
    degrade(make_text_image()).save(p, quality=35)
    return p


@pytest.fixture
def upside_down_png(tmp_path):
    p = tmp_path / "flipped.png"
    make_text_image().rotate(180).save(p)
    return p


@pytest.fixture
def multi_tiff(tmp_path):
    p = tmp_path / "multi.tiff"
    a = make_text_image()
    b = make_text_image(["Second frame", "Invoice No: INV-2024-0099", "Total: 77.00 EUR"])
    a.save(p, save_all=True, append_images=[b])
    return p


@pytest.fixture
def multi_sheet_xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["ACME Corp - Q1 Invoice Register"])  # title row above the header
    ws.append([])
    ws.append(["Invoice No", "Date", "Supplier", "Total"])
    ws.append(["00123", "2024-01-05", "Acme Trading", 1150.0])
    ws.append(["00124", "2024-01-09", "Globex", 99.5])
    ws2 = wb.create_sheet("Lines")
    ws2.append(["Invoice No", "Item", "Qty", "Price"])
    ws2.append(["00123", "Widget", 10, 100.0])
    wb.create_sheet("Empty")
    p = tmp_path / "book.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def semicolon_csv(tmp_path):
    p = tmp_path / "eu.csv"
    p.write_bytes("Rechnung;Lieferant;Betrag\n00123;Müller GmbH;1150,50\n00124;Société Générale;99,00\n".encode("cp1252"))
    return p


@pytest.fixture
def real_xls(tmp_path):
    """A genuine legacy .xls binary (written by xlwt), not a fake header."""
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Invoices")
    ws.write(0, 0, "ACME Corp - Legacy Register")
    for c, h in enumerate(["Invoice No", "Date", "Supplier", "Total"]):
        ws.write(2, c, h)
    for r, row in enumerate([["00123", "2024-01-05", "Acme Trading", 1150.0],
                             ["00124", "2024-01-09", "Globex", 99.5]], start=3):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    wb.add_sheet("Second").write(0, 0, "Item")
    p = tmp_path / "legacy.xls"
    wb.save(str(p))
    return p
